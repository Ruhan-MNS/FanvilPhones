# -*- coding: utf-8 -*-
"""Generates the FancilPhones Back-End Documentation (.docx),
styled after the Multi-Net Systems 'Tender Hub Documentation'."""
import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RED = RGBColor(0xE2, 0x23, 0x1A)      # the app's brand red (#E2231A)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)
CODE_BG = "F4F4F4"
HDR_BG = "E2231A"

doc = Document()

# ---- base styles ----
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)

for lvl, sz in [('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)]:
    st = doc.styles[lvl]
    st.font.name = 'Calibri'
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = RED if lvl == 'Heading 1' else DARK


def shade(cell_or_par_xml, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    cell_or_par_xml.append(shd)


def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def h3(t): doc.add_heading(t, level=3)


def para(t="", bold=False, italic=False, color=None, size=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p


def bullet(t, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.add_run(t)
    return p


def numbered(t):
    p = doc.add_paragraph(style='List Number')
    p.add_run(t)
    return p


def code(text):
    """Mono, shaded code block."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shade(pPr, CODE_BG)
    # thin border
    borders = OxmlElement('w:pBdr')
    for edge in ('top', 'bottom', 'left', 'right'):
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '6'); b.set(qn('w:color'), 'DDDDDD')
        borders.append(b)
    pPr.append(borders)
    for i, line in enumerate(text.split('\n')):
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        r.font.color.rgb = DARK
        if i < len(text.split('\n')) - 1:
            r.add_break()
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(htext)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)
        shade(hdr[i]._tc.get_or_add_tcPr(), HDR_BG)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def toc():
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    it = OxmlElement('w:r'); itt = OxmlElement('w:t')
    itt.text = "Right-click and choose 'Update Field' to build the table of contents."
    it.append(itt); fld.append(it)
    p._p.append(fld)


def hr():
    para("")


# ============================================================
# TITLE PAGE
# ============================================================
sp = doc.add_paragraph()
sp.paragraph_format.space_before = Pt(120)
para("FANCIL PHONES", bold=True, color=RED, size=40, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Back-End Documentation", bold=True, color=DARK, size=24, align=WD_ALIGN_PARAGRAPH.CENTER)
para("", )
para("Fleet Phonebook Synchronisation for Fanvil IP Phones", italic=True,
     color=GREY, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(6):
    para("")
para("Blazor Server  ·  .NET 10  ·  MudBlazor  ·  EF Core (SQLite)",
     color=GREY, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
para("")
para("Author: Ruhan Slater", color=DARK, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Revision 01  ·  " + datetime.date.today().strftime("%B %Y"),
     color=GREY, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ============================================================
# CONTENTS
# ============================================================
h1("Contents")
toc()
doc.add_page_break()

# ============================================================
# 1. INTRODUCTION
# ============================================================
h1("Introduction")
para("FancilPhones is an in-house Blazor Server application used to maintain a single master "
     "contact list and push it out to a fleet of Fanvil X-series IP phones. Each handset stores "
     "its own local phonebook that can normally only be edited through that phone's individual web "
     "admin page, which makes keeping a large fleet consistent slow and error-prone. FancilPhones "
     "centralises that editing: contacts are managed once, and the application logs in to each phone "
     "over the network and replaces its phonebook automatically.")
para("The application also discovers phones on the local network, pushes SIP line settings "
     "(display name and extension), records every action against each phone, and exports/imports "
     "both the contact list and the fleet configuration as CSV and Excel files.")

h2("Technology Stack")
table(
    ["Concern", "Technology"],
    [
        ["Runtime / framework", "ASP.NET Core on .NET 10"],
        ["UI", "Blazor Server (Interactive Server render mode)"],
        ["Component library", "MudBlazor 9.4"],
        ["Data access", "Entity Framework Core 10 (SQLite provider)"],
        ["Database", "SQLite — single file fancilphones.db"],
        ["Authentication", "ASP.NET Core Identity (cookie based)"],
        ["Spreadsheet export/import", "ClosedXML 0.104"],
        ["Phone integration", "Custom raw-socket HTTP client (FanvilHttpClient)"],
    ],
    widths=[2.6, 4.4])

# ============================================================
# 2. ARCHITECTURE
# ============================================================
h1("Architecture")
para("Unlike a multi-project Onion solution, FancilPhones is a single ASP.NET Core project organised "
     "into clear internal layers by folder. The application is small and self-contained, so the layers "
     "are namespaces/folders within one project rather than separate class libraries. The layers are:")
numbered("Presentation — Razor components and pages (the Blazor UI).")
numbered("Services — the business logic: phone communication, sync, discovery, import/export.")
numbered("Data — the EF Core entities and the database context.")
numbered("Host / Composition — Program.cs, which wires everything together and exposes the minimal-API endpoints.")

h2("Breakdown of the Layers")
h3("1. Presentation")
para("The Presentation layer lives under Components/. It contains the Blazor pages the user "
     "interacts with (Home, Phones, Contacts, Sync, Logs, Users, Login) plus shared layout and "
     "dialog components. Pages do not talk to the phones directly — they inject the services in the "
     "Services layer and the EF Core context factory for data.")
h3("2. Services")
para("The Services layer (Services/) holds the bulk of the application's functionality: the lenient "
     "HTTP client that speaks to Fanvil phones, the synchronisation service, the network discovery "
     "service, and the CSV/XLSX helpers. Services read and write the database through the EF Core "
     "context factory and never depend on UI code.")
h3("3. Data")
para("The Data layer (Data/) contains the entity models (Phone, Contact, SyncRun, AppUser) and the "
     "AppDbContext that maps them to the SQLite database. It is the single point through which the "
     "rest of the application reaches the database.")
h3("4. Host / Composition")
para("Program.cs is the composition root. It registers services in the dependency-injection "
     "container, configures Identity and authentication, ensures the database exists and is seeded, "
     "and maps the minimal-API endpoints used for file downloads and authentication form posts.")

# ============================================================
# 3. FILE STRUCTURE
# ============================================================
h1("File Structure")
para("The solution FancilPhones.sln contains a single web project, FancilPhones. Its important "
     "folders and files are described below.")

h2("Project ‘FancilPhones’")
h3("Dependencies (NuGet packages)")
table(
    ["Package", "Version", "Purpose"],
    [
        ["Microsoft.EntityFrameworkCore.Sqlite", "10.0.8", "EF Core SQLite database provider"],
        ["Microsoft.EntityFrameworkCore.Design", "10.0.8", "Design-time EF Core tooling"],
        ["Microsoft.AspNetCore.Identity.EntityFrameworkCore", "10.0.8", "Identity user/role storage"],
        ["MudBlazor", "9.4.0", "Blazor UI component library"],
        ["ClosedXML", "0.104.2", "Reading/writing .xlsx workbooks"],
    ],
    widths=[3.4, 1.0, 2.6])

h3("Components/")
para("All front-end Razor code.")
bullet("App.razor, Routes.razor — the application root and router.")
bullet("Layout/ — MainLayout, EmptyLayout, NavMenu, ReconnectModal and their CSS.")
bullet("Pages/ — the routable pages and dialog components (see the Pages and Components section).")
bullet("_Imports.razor — global using directives for components.")

h3("Data/")
para("The EF Core entities and database context: AppDbContext.cs, AppUser.cs, Phone.cs, "
     "Contact.cs, SyncRun.cs.")

h3("Services/")
para("The service classes: FanvilHttpClient.cs, PhoneSyncService.cs, PhoneDiscoveryService.cs, "
     "PhonebookCsv.cs, PhonesXlsx.cs.")

h3("Other files")
bullet("Program.cs — composition root, DI registration, endpoints, database seeding.")
bullet("appsettings.json / appsettings.Development.json — logging and connection-string configuration.")
bullet("fancilphones.db — the SQLite database file (created automatically on first run).")
bullet("Properties/launchSettings.json — local launch profiles.")

# ============================================================
# 4. DATABASE
# ============================================================
h1("Database")
para("FancilPhones uses a single SQLite database, fancilphones.db, created in the project directory. "
     "The connection string defaults to Data Source=fancilphones.db and can be overridden by the "
     "ConnectionStrings:Default setting.")
para("The schema is created with EF Core's EnsureCreated() (there are no migrations). Because "
     "EnsureCreated does not alter an existing schema, Program.cs performs a small set of additive "
     "ALTER TABLE statements at start-up — guarded by a PRAGMA table_info check so columns are only "
     "added when missing. This lets the schema evolve on already-deployed databases without "
     "migrations.")

h2("Tables")
table(
    ["Table", "Stores"],
    [
        ["Phones", "Each Fanvil handset: name, IP, credentials, upload settings, SIP line settings, last-sync status."],
        ["Contacts", "The master phonebook entries that get pushed to every phone."],
        ["SyncRuns", "A history row for every Sync, Push or Probe action against a phone."],
        ["AspNetUsers / AspNetRoles / …", "Standard ASP.NET Identity tables for users and roles."],
    ],
    widths=[2.2, 4.8])

h2("Indexes and Relationships")
bullet("Phones.IpAddress has a unique index — no two phones can share an IP.")
bullet("Contacts.DisplayName is indexed for ordered listing.")
bullet("SyncRun has a required foreign key to Phone with cascade delete — deleting a phone removes its history.")

h2("Additive Columns Applied at Start-up")
para("The following columns are added on existing databases if absent (see Program.cs):")
code('AddColumnIfMissing("Phones",   "SipDisplayName",         "TEXT NULL");\n'
     'AddColumnIfMissing("Phones",   "SipLineIndex",           "INTEGER NOT NULL DEFAULT 1");\n'
     'AddColumnIfMissing("Phones",   "SipExtension",           "TEXT NULL");\n'
     'AddColumnIfMissing("Phones",   "SipRegistrationEnabled", "INTEGER NOT NULL DEFAULT 1");\n'
     'AddColumnIfMissing("SyncRuns", "Action",                 "TEXT NOT NULL DEFAULT \'Sync\'");')

# ============================================================
# 5. AUTHENTICATION
# ============================================================
h1("Authentication")
para("There are two distinct kinds of authentication in FancilPhones: application sign-in (how a "
     "user logs in to FancilPhones itself) and phone authentication (how the application logs in to "
     "a Fanvil handset). They are unrelated and use completely different mechanisms.")

h2("Application Sign-In")
para("Application users are managed with ASP.NET Core Identity, configured in Program.cs with a "
     "cookie scheme. Key points:")
bullet("Identity is registered with AddIdentityCore, AddRoles and AddEntityFrameworkStores; sign-in uses a 30-day sliding cookie.")
bullet("The login form posts to the /auth/login minimal-API endpoint, which signs the user in and redirects.")
bullet("Password policy is deliberately relaxed (minimum length 4, no complexity) because the tool runs on a trusted internal network. This should be tightened if the app is ever exposed publicly.")
bullet("On first run an admin account is seeded (username admin, password admin) with MustChangePassword set, forcing a password change on first login.")
bullet("AppUser extends IdentityUser with a single extra flag, MustChangePassword.")

h2("Roles")
para("Two roles are seeded at start-up. Pages are restricted with the [Authorize] attribute; the "
     "Users page additionally requires the Admin role.")
table(
    ["Role", "Access"],
    [
        ["Admin", "Full access, including the Users page (create users, reset passwords, assign roles)."],
        ["Technician", "Access to all operational pages (Phones, Contacts, Sync, Logs) but not user management."],
    ],
    widths=[1.8, 5.2])

h2("Phone Authentication (Fanvil nonce / MD5 challenge)")
para("Fanvil X-series phones do NOT support HTTP Basic authentication — sending an Authorization "
     "header causes the phone's embedded ‘Rapid Logic’ web server to reset the connection. Instead "
     "they use a nonce + MD5 challenge that is bound to a single TCP connection. The full sequence "
     "is implemented in FanvilHttpClient.LoginAsync and is described in the Fanvil Integration "
     "appendix.")

# ============================================================
# 6. FUNCTIONALITY
# ============================================================
h1("Functionality")

# ---- Data Context ----
h2("Data Context")
para("AppDbContext is the application's database context. It derives from IdentityDbContext so it "
     "owns both the application tables and the Identity tables, and exposes three DbSet properties: "
     "Phones, Contacts and SyncRuns.")
para("Two registrations are made in Program.cs: a DbContextFactory (used by Blazor pages, which are "
     "long-lived and need short-lived contexts) and a scoped AppDbContext (used by Identity). "
     "OnModelCreating configures the unique IP index, the DisplayName index and the SyncRun→Phone "
     "relationship.")
code('public class AppDbContext : IdentityDbContext<AppUser, IdentityRole, string>\n'
     '{\n'
     '    public DbSet<Phone> Phones => Set<Phone>();\n'
     '    public DbSet<Contact> Contacts => Set<Contact>();\n'
     '    public DbSet<SyncRun> SyncRuns => Set<SyncRun>();\n'
     '\n'
     '    protected override void OnModelCreating(ModelBuilder b)\n'
     '    {\n'
     '        base.OnModelCreating(b);\n'
     '        b.Entity<Phone>().HasIndex(p => p.IpAddress).IsUnique();\n'
     '        b.Entity<Contact>().HasIndex(c => c.DisplayName);\n'
     '        b.Entity<SyncRun>().HasOne(r => r.Phone).WithMany()\n'
     '            .HasForeignKey(r => r.PhoneId).OnDelete(DeleteBehavior.Cascade);\n'
     '    }\n'
     '}')

# ---- Models ----
h2("Models")
para("The models are plain C# classes mapped to database tables. Each is described below.")

h3("Phone.cs")
para("Represents one Fanvil handset and everything needed to reach and configure it.")
table(
    ["Property", "Type", "Description"],
    [
        ["Id", "int", "Primary key."],
        ["Name", "string", "Friendly name (e.g. ‘Reception’)."],
        ["IpAddress", "string", "The phone's IP; unique across the fleet."],
        ["Username / Password", "string", "Web-admin credentials (default admin/admin)."],
        ["Scheme", "string", "http or https."],
        ["UploadPath", "string", "Web path that receives the phonebook upload (default /pBookAdv.htm)."],
        ["UploadFieldName", "string", "Multipart field name for the file (default PHONEBOOK)."],
        ["Enabled", "bool", "Whether ‘Sync all’ includes this phone."],
        ["SipDisplayName", "string?", "Optional SIP display name pushed to the line settings."],
        ["SipLineIndex", "int", "Which SIP line the push targets (1-based)."],
        ["SipExtension", "string?", "Optional SIP extension / username."],
        ["SipRegistrationEnabled", "bool", "Whether the push enables SIP registration on the line."],
        ["LastSyncedAt", "DateTime?", "Timestamp of the last action."],
        ["LastSyncStatus / LastSyncMessage", "string?", "Outcome of the last action."],
    ],
    widths=[2.3, 1.1, 3.6])

h3("Contact.cs")
para("One entry in the master phonebook. These rows are serialised to the CSV that is uploaded to "
     "every phone.")
table(
    ["Property", "Type", "Description"],
    [
        ["Id", "int", "Primary key."],
        ["DisplayName", "string", "Contact name (required, indexed)."],
        ["OfficeNumber", "string?", "‘work’ number in the Fanvil CSV."],
        ["MobileNumber", "string?", "‘mobile’ number."],
        ["OtherNumber", "string?", "‘other’ number."],
        ["Group", "string?", "Phonebook group."],
        ["Ring", "string?", "Ringtone (default ‘Auto’)."],
        ["UpdatedAt", "DateTime", "Last edit time (UTC)."],
    ],
    widths=[2.0, 1.0, 4.0])

h3("SyncRun.cs")
para("An audit row written for every action against a phone. Powers the Logs page.")
table(
    ["Property", "Type", "Description"],
    [
        ["Id", "int", "Primary key."],
        ["PhoneId / Phone", "int", "The phone this run targeted (FK, cascade delete)."],
        ["StartedAt / FinishedAt", "DateTime", "Run timing."],
        ["Status", "string", "Pending, Running, Success, Failed or Error."],
        ["Action", "string", "Sync (phonebook), Push (SIP line) or Probe (reachability)."],
        ["HttpStatusCode", "int?", "HTTP status returned by the phone."],
        ["Message", "string?", "Human-readable outcome."],
        ["ContactCount", "int", "Number of contacts in the pushed phonebook."],
    ],
    widths=[2.2, 1.0, 3.8])

h3("AppUser.cs")
para("Extends the standard Identity user with one field, MustChangePassword, which forces a "
     "password change after the account is seeded or reset by an administrator.")

# ---- Services ----
h2("Services")
para("The Services layer contains all of the application's behaviour. Each service is described "
     "below with its key methods.")

h3("FanvilHttpClient.cs")
para("A deliberately lenient HTTP/1.1 client written specifically for Fanvil phones. The framework's "
     "HttpClient cannot be used because the phone's embedded server emits non-RFC-compliant responses "
     "(it frequently resets the socket after the body instead of closing cleanly) and ties the login "
     "nonce to a single TCP connection. FanvilHttpClient therefore talks over a raw TcpClient socket, "
     "keeps ONE persistent connection alive for an entire login/upload sequence, and parses responses "
     "tolerantly.")
para("Key members:")
bullet("LoginAsync — performs the nonce/MD5 challenge login (see the Fanvil Integration appendix). Retries up to 4 times on a fresh connection because the nonce endpoint is flaky.")
bullet("LogoutAsync — posts the Logout form so the phone frees its (limited) session slot. Best-effort.")
bullet("DeleteAllContactsAsync — mirrors the web UI's ‘Delete All’ button so a sync becomes a true mirror.")
bullet("PostFormAsync / PostFileAsync — url-encoded form posts and multipart file uploads.")
bullet("GetAsync — a GET request, used for nonce retrieval and form scraping.")
para("Internally it manages cookies across requests, builds raw HTTP requests by hand (including a "
     "matching Referer header, which some firmwares require), and reads exactly one response bounded "
     "by Content-Length or a short idle gap.")

h3("PhoneSyncService.cs")
para("The orchestrator that drives FanvilHttpClient to perform the three phone actions and records a "
     "SyncRun for each. Registered as a scoped service and injected into the Phones and Sync pages.")
para("Public methods:")
table(
    ["Method", "Purpose"],
    [
        ["SyncAllAsync", "Builds the CSV once and pushes it to every enabled phone in parallel."],
        ["SyncOneAsync", "Pushes the phonebook to a single phone by id."],
        ["PushSipDisplayNameAsync", "Pushes the SIP display name / extension to the phone's /lines.htm line settings."],
        ["ProbeAsync", "Tests reachability by logging in and straight back out, freeing the session slot."],
    ],
    widths=[2.6, 4.4])
para("The phonebook sync sequence (PushAsync) is: log in → delete all existing contacts → upload the "
     "master CSV → log out. The delete-all step is required because the phone's CSV import only merges "
     "(it never removes), so without it deletions in the app would never propagate. The status and "
     "message are written both to a SyncRun row and to the phone's Last-sync fields.")
para("PushSipDisplayNameAsync is more delicate: the phone's /lines.htm form echoes saved passwords as "
     "masked strings and is populated by JavaScript, so blindly re-posting every field would wipe the "
     "SIP configuration. The method therefore submits only the fields it intends to change plus the "
     "minimum state Fanvil needs to commit, and explicitly skips any password-like field.")

h3("PhoneDiscoveryService.cs")
para("Scans an IP range to find Fanvil phones on the network, used by the Scan Network dialog.")
bullet("EnumerateIps — expands a user-entered IPv4 address, CIDR subnet or hyphen-separated range into individual addresses, excluding network/broadcast for /24 and smaller. Discovery requires the operator to enter the intended range.")
bullet("ProbeOneAsync — cheaply TCP-connects to port 80, then fetches / and /login.htm and checks for Fanvil signatures (e.g. logonButton, RapidLogic, frame markers). Deliberately avoids the nonce endpoint so scanning doesn't burn session slots on live handsets.")
bullet("ScanAsync — streams results as each probe completes, using a bounded-concurrency worker pool (default 32) and a channel so the UI can show progress live.")

h3("PhonebookCsv.cs")
para("Builds and parses the exact CSV format Fanvil phones import and export. The header is "
     '"name","work","mobile","other","ring","groups" with every field double-quoted and CRLF line '
     "endings.")
bullet("Build — serialises the master contacts into that CSV (UTF-8, no BOM). An empty Ring becomes ‘Default’.")
bullet("Parse — reads CSV text back into Contact records, mapping columns by header name (any order) and tolerating a missing header by falling back to fixed positions. Used by the Contacts page's Import CSV feature.")

h3("PhonesXlsx.cs")
para("Parses a Phones export workbook back into Phone objects, used to bulk-import or restore the "
     "fleet configuration. Column lookups are case-insensitive and accept several header aliases "
     "(e.g. IP, IpAddress, ‘IP Address’); missing optional columns fall back to the Phone model's "
     "defaults. Only Name and IP are mandatory.")

# ---- Endpoints ----
h2("Minimal-API Endpoints")
para("In place of MVC controllers, FancilPhones exposes a handful of minimal-API endpoints mapped in "
     "Program.cs. They cover file downloads and the authentication form posts.")
table(
    ["Endpoint", "Method", "Auth", "Description"],
    [
        ["/api/phonebook.csv", "GET", "Open", "Downloads the master contact list as the Fanvil CSV."],
        ["/api/phonebook.xlsx", "GET", "Open", "Downloads the contacts as a formatted Excel workbook."],
        ["/api/phones.xlsx", "GET", "Required", "Exports the fleet configuration (backup/restore)."],
        ["/auth/login", "POST", "Open", "Validates credentials, signs in, handles forced password change."],
        ["/auth/logout", "POST", "Open", "Signs the user out and redirects to /login."],
    ],
    widths=[1.9, 0.7, 0.9, 3.5])
para("The phonebook CSV/XLSX downloads are intentionally left open (they contain only the shared "
     "contact list, no secrets), while the phones export — which contains credentials — requires "
     "authentication. The auth endpoints disable antiforgery so the static-rendered login form can "
     "post to them.")

# ---- Pages ----
h2("Pages and Components")
para("The Razor pages live under Components/Pages/. All operational pages are marked [Authorize]; the "
     "Users page additionally requires the Admin role. Pages inject the EF Core context factory and "
     "the relevant services rather than accessing phones directly.")
table(
    ["Page / Route", "Purpose"],
    [
        ["Home — /", "Dashboard / landing page with fleet overview."],
        ["Phones — /phones", "Manage handsets: add/edit, probe, sync one, push SIP settings, scan network, import/export."],
        ["Contacts — /contacts", "Manage the master phonebook: add/edit/delete contacts, import CSV, export."],
        ["Sync — /sync", "Run ‘Sync all’ and watch per-phone results."],
        ["Logs — /logs", "Browse the SyncRun history (Sync / Push / Probe outcomes)."],
        ["Users — /users", "Admin only: create users, reset passwords, assign roles."],
        ["Login — /login", "Sign-in form (posts to /auth/login)."],
        ["ChangePassword — /account/change-password", "Forced/voluntary password change."],
    ],
    widths=[2.6, 4.4])
para("Supporting dialog components include ContactEditDialog, PhoneEditDialog, CreateUserDialog, "
     "ResetPasswordDialog and ScanNetworkDialog. Layout components (MainLayout, NavMenu, "
     "ReconnectModal) provide the shell and navigation.")

# ============================================================
# APPENDIX — Fanvil integration
# ============================================================
doc.add_page_break()
h1("Appendix A — Fanvil Phone Integration")
para("This appendix documents the hard-won details of communicating with Fanvil X-series phones. "
     "These behaviours are non-standard and are the main reason a custom HTTP client exists.")

h2("Login Sequence (nonce / MD5 challenge)")
para("All steps run over ONE persistent TCP connection — the nonce is connection-bound.")
numbered("GET / to open the session.")
numbered("GET /key==nonce?now=<unix-ms> — the phone returns a short (~16-char) nonce in the body.")
numbered("Store the nonce as the auth cookie.")
numbered("Compute encoded = username + \":\" + md5(username + \":\" + password + \":\" + nonce).")
numbered("POST / with form fields encoded, CurLanguage=en, ReturnPage=/. A login failure re-serves the login page (body contains logonButton).")
code('var nonceResp = await GetAsync($"/key==nonce?now={now}");\n'
     'var nonce = nonceResp.Body.Trim();\n'
     '_cookies["auth"] = nonce;\n'
     'var digest  = Md5Hex($"{username}:{password}:{nonce}");\n'
     'var encoded = $"{username}:{digest}";\n'
     'await PostFormAsync("/", new() {\n'
     '    ["encoded"] = encoded, ["CurLanguage"] = "en", ["ReturnPage"] = "/" });')

h2("Session Slots")
para("Each phone has a LIMITED number of session slots. Every successful login consumes one; if it "
     "is not released, the nonce endpoint starts returning empty bodies (Content-Length: 0). The "
     "application therefore always logs out after every action (POST /title.htm with "
     "DefaultLogout=Logout). If nonces are already failing for everything, rebooting the phone clears "
     "the stale sessions.")

h2("Phonebook Upload")
para("The phonebook is uploaded as multipart form-data to /pBookAdv.htm under the field name "
     "PHONEBOOK. Because the phone's import only merges, a true mirror requires deleting all existing "
     "contacts first (POST /contacts.htm with DefaultDeleteAll=Delete All). The application uploads "
     "CSV rather than XML.")

h2("SIP Line Settings")
para("SIP display name and extension are written to /lines.htm (fields SIP_DisPlayName_R, "
     "SIP_PhoneNum_R and SIP_RegUser_R; the ‘Activate’ checkbox is SIP_EnableSipReg_RW). The form is "
     "JavaScript-populated and echoes passwords as masked strings, so only the changed fields plus a "
     "minimal set of state fields (DefaultSubmit=Apply, CheckBoxManager, keepList, "
     "SIP_PhoneLineEntry) are submitted; password-like fields are never posted back.")

h2("Network Quirks Summary")
bullet("No HTTP Basic auth — an Authorization header makes the server reset the connection.")
bullet("Responses are often terminated by a socket reset rather than a clean close; the parser treats post-body read errors as end-of-stream.")
bullet("The web UI is frame-based — the real login form lives in /login.htm, not /.")
bullet("Some firmwares reject form posts whose Referer header doesn't match the form page.")

doc.save(r"C:\Users\ruhan\RiderProjects\FancilPhones\FancilPhones Backend Documentation (Rev.01).docx")
print("saved")
